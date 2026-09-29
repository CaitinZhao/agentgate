"""Excel case-bank workflow (v2): template generation, validation (auto-completion + red marking), bank import.

v2 columns mirror the gold skeleton (doc 35 §4.1): the author fills query/type/gold
(value or answer + optional aliases/keywords), optional checkpoints and rubric —
pack-level config (evidence/caliber/red lines) is inherited, never typed per row.
"""
from typing import Dict, List, Tuple

from .models import Case, CaseSource, DiagnosisHint, Gold, Checkpoint, RubricPoint

COLUMNS = [
    "case_id", "suite", "level", "status", "as_of",
    "origin", "seed", "adaptation",
    "query", "type(auto/numeric/boolean/extractive/free_text/refusal)",
    "gold_value", "gold_unit", "gold_tol_rel",
    "gold_aliases(多值用；)", "gold_keywords(多值用；)",
    "checkpoints(多值用；格式 signal|pattern|desc|weight)",
    "rubric(多值用；格式 要点|权重)",
    "failure_kind", "target_layer", "expected_behavior",
]

LEVEL_OPTIONS = "L0,L1,L2"
BOOL_OPTIONS = "yes,no"
TYPE_OPTIONS = "auto,numeric,boolean,extractive,free_text,refusal"
_LEVELS = {"L0", "L1", "L2"}
_LEVEL_DEFAULT = "L2"
NL = chr(10)
SEP = "；"

_BAD_FILL = "FFFFC7CE"      # light red fill (Excel "bad" style)
_BAD_FONT = "FF9C0006"      # dark red font


def _split(v) -> List[str]:
    if v is None:
        return []
    parts = str(v).replace("；", ";").replace(",", ";").split(";")
    return [x.strip() for x in parts if x.strip()]


def _bool(v) -> bool:
    return str(v).strip().lower() in ("yes", "y", "true", "1", "是")


def _gold_from_case(d: Dict) -> Dict:
    g = d.get("gold", {}) or {}
    f = g.get("final", {}) or {}
    return {"final": f, "checkpoints": g.get("checkpoints", []), "rubric": g.get("rubric", [])}


def _row_from_case(d: Dict) -> List:
    gold = _gold_from_case(d)
    f = gold["final"]
    cps = SEP.join("%s|%s|%s|%s" % (c.get("signal", "content"), c.get("pattern", ""),
                                    c.get("desc", ""), c.get("weight", 1))
                   for c in gold["checkpoints"])
    rubric = SEP.join("%s|%s" % (r.get("point", ""), r.get("weight", 1))
                      for r in gold["rubric"])
    return [
        d.get("case_id", ""), d.get("suite", ""), d.get("level", "L2"),
        d.get("status", "active"), d.get("as_of") or "",
        d.get("source", {}).get("origin", ""), d.get("source", {}).get("seed", ""),
        d.get("source", {}).get("adaptation", ""),
        d.get("input", {}).get("query", ""), d.get("type", "auto"),
        f.get("value") if not isinstance(f.get("value"), (dict, list)) else "",
        f.get("unit") or "", f.get("tol_rel", 0.01),
        SEP.join(str(a) for a in f.get("aliases", [])),
        SEP.join(str(k) for k in f.get("keywords", [])),
        cps, rubric,
        d.get("diagnosis_hint", {}).get("failure_kind", ""),
        d.get("diagnosis_hint", {}).get("target_layer", ""),
        d.get("diagnosis_hint", {}).get("expected_behavior", ""),
    ]


def _add_dropdowns(ws):
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.utils import get_column_letter

    def dropdown(col_name, options):
        idx = COLUMNS.index(col_name) + 1
        col = get_column_letter(idx)
        dv = DataValidation(type="list", formula1='"%s"' % options, allow_blank=True)
        ws.add_data_validation(dv)
        dv.add("%s2:%s1000" % (col, col))

    dropdown("level", LEVEL_OPTIONS)
    dropdown("status", "active,retired")
    dropdown(COLUMNS[9], TYPE_OPTIONS)


def _style(ws, notes=None):
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
    for c in ws[1]:
        c.font = Font(bold=True)
    for i, name in enumerate(COLUMNS, 1):
        ws.column_dimensions[get_column_letter(i)].width = max(12, min(42, len(name) + 8))
    ws.freeze_panes = "A2"
    _add_dropdowns(ws)
    if notes:
        ws2 = ws.parent.create_sheet("说明")
        for i, t in enumerate(notes, 1):
            ws2.cell(row=i, column=1, value=t)


def export_template(out_path: str) -> str:
    """Generate a blank Excel template: header + dropdowns + one sample row (replace or delete when filling)."""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "cases"
    ws.append(COLUMNS)
    sample = _row_from_case({
        "case_id": "sample-001（示例行，替换或删除）",
        "suite": "your-suite", "level": "L2",
        "input": {"query": "在这里写你要问 Agent 的问题"},
        "type": "numeric",
        "gold": {"final": {"value": 42, "unit": "%", "tol_rel": 0.01},
                 "checkpoints": [{"signal": "tool", "pattern": "retrieve_report",
                                  "desc": "调用检索工具", "weight": 1}]},
    })
    ws.append(sample)
    _style(ws, notes=[
        "level：L0 核心 / L1 领域 / L2 全量，默认 L2",
        "type：留 auto 由判定系统按 gold 形状推断；数值题填 numeric 并填 gold_value/unit/tol_rel",
        "是非题：type=boolean，gold_value 填 yes/no；抽取题填 extractive + gold_value（可加 gold_aliases）",
        "refusal 题：gold_keywords 填必须出现的拒答关键词（多个用；分隔）",
        "checkpoints 每项格式：signal|pattern|desc|weight，signal 取 tool/content，多项用；分隔",
        "rubric（free_text 题的评分要点）每项格式：要点|权重，多项用；分隔",
        "领域配置（证据/口径/红线）在题库的领域包里配置，出题人不填",
        "改完后运行：agentgate xlsx-import --path 本文件 --db <题库路径>/cases.db",
    ])
    wb.save(out_path)
    return out_path


def export_cases(cases, out_path: str) -> str:
    """Existing bank -> Excel (for bulk editing)."""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "cases"
    ws.append(COLUMNS)
    for c in cases:
        ws.append(_row_from_case(c.model_dump()))
    _style(ws)
    wb.save(out_path)
    return out_path


def _parse_checkpoints(v: str) -> List[Checkpoint]:
    out = []
    for part in _split(v):
        seg = [s.strip() for s in part.replace("｜", "|").split("|")]
        signal = seg[0] if seg and seg[0] in ("tool", "content", "state") else "content"
        pattern = seg[1] if len(seg) > 1 else ""
        desc = seg[2] if len(seg) > 2 else pattern
        try:
            weight = float(seg[3]) if len(seg) > 3 else 1.0
        except ValueError:
            weight = 1.0
        if pattern:
            out.append(Checkpoint(desc=desc, signal=signal, pattern=pattern, weight=weight))
    return out


def _parse_rubric(v: str) -> List[RubricPoint]:
    out = []
    for part in _split(v):
        seg = [s.strip() for s in part.replace("｜", "|").split("|")]
        if not seg or not seg[0]:
            continue
        try:
            weight = float(seg[1]) if len(seg) > 1 else 1.0
        except ValueError:
            weight = 1.0
        out.append(RubricPoint(point=seg[0], weight=weight))
    return out


def import_xlsx(path: str, suite_fallback: str = "", import_date: str = None,
                auto_fix: bool = True) -> Tuple[List[Case], List[Dict], List[str]]:
    """Excel -> bank: validate row by row; auto-complete what is possible; returns (valid Cases, problems, auto-fix notes).

    problems: [{row: row number, msg: error}] - the caller uses it to generate the red-marked Excel.
    """
    from openpyxl import load_workbook
    ws = load_workbook(path, data_only=True).active
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    cases: List[Case] = []
    problems: List[Dict] = []
    fixes: List[str] = []
    seen: Dict[str, int] = {}
    import re as _re
    id_pat = _re.compile("^[a-z][a-z0-9-]*$")

    for i, row in enumerate(rows, start=2):
        if not row or not any(v not in (None, "") for v in row):
            continue
        d = dict(zip(COLUMNS, row))
        row_errors: List[str] = []

        cid = str(d.get("case_id") or "").strip()
        suite = str(d.get("suite") or "").strip() or suite_fallback
        if not cid:
            seq = sum(1 for c in cases if (c.suite or suite) == suite) + 1
            cid = "%s-%03d" % ((suite or "case").lower(), seq)
            fixes.append("第%d行：case_id 留空，自动生成 %s" % (i, cid))
        elif not id_pat.match(cid):
            row_errors.append("case_id 格式不对（%r）：只能用小写字母/数字/连字符，且以字母开头" % cid)

        if cid in seen:
            row_errors.append("case_id 与第%d行重复" % seen[cid])
        else:
            seen[cid] = i

        query = str(d.get("query") or "").strip()
        if not query:
            row_errors.append("query（题目内容）不能为空")

        level = str(d.get("level") or _LEVEL_DEFAULT).strip()
        if not level:
            level = _LEVEL_DEFAULT
            fixes.append("第%d行：level 留空，默认 L2" % i)
        elif level not in _LEVELS:
            row_errors.append("level（%s）不在可选范围 %s" % (level, sorted(_LEVELS)))

        as_of = str(d.get("as_of") or "").strip() or (import_date or "")
        if d.get("as_of") in (None, "") and import_date:
            fixes.append("第%d行：as_of 留空，默认填入池日期 %s" % (i, import_date))

        ctype = str(d.get(COLUMNS[9]) or "auto").strip().lower()
        if ctype not in ("auto", "numeric", "boolean", "extractive", "free_text", "refusal"):
            row_errors.append("type（%s）不在可选范围 %s" % (ctype, TYPE_OPTIONS))

        gvalue = d.get("gold_value")
        gold_final: Dict = {}
        if gvalue not in (None, ""):
            if ctype == "numeric" or (ctype == "auto" and isinstance(gvalue, (int, float))):
                try:
                    gold_final = {"value": float(gvalue),
                                  "unit": str(d.get("gold_unit") or "") or "none",
                                  "tol_rel": float(d.get("gold_tol_rel") or 0.01)}
                except (TypeError, ValueError):
                    row_errors.append("gold_value（%r）不是数字（numeric 题）" % gvalue)
            else:
                gold_final = {"value": str(gvalue).strip(),
                              "aliases": _split(d.get("gold_aliases(多值用；)"))}
        keywords = _split(d.get("gold_keywords(多值用；)"))
        if keywords:
            gold_final["keywords"] = keywords
        if ctype == "refusal" and not keywords:
            fixes.append("第%d行：refusal 题建议填 gold_keywords（拒答关键词）" % i)
        if ctype == "free_text" and not _split(d.get("rubric(多值用；格式 要点|权重)")):
            fixes.append("第%d行：free_text 题建议填 rubric（评分要点），否则只能落人工复核" % i)

        checkpoints = _parse_checkpoints(d.get("checkpoints(多值用；格式 signal|pattern|desc|weight)"))
        rubric = _parse_rubric(d.get("rubric(多值用；格式 要点|权重)"))

        if row_errors:
            for msg in row_errors:
                problems.append({"row": i, "msg": msg})
            continue

        cases.append(Case(
            case_id=cid, suite=suite or None, level=level,
            status=str(d.get("status") or "active"),
            as_of=as_of or None,
            source=CaseSource(origin=str(d.get("origin") or "manual"),
                              seed=str(d.get("seed") or "") or None,
                              adaptation=str(d.get("adaptation") or "") or None),
            input={"query": query},
            type=ctype,
            gold=Gold(final=gold_final, checkpoints=checkpoints, rubric=rubric),
            diagnosis_hint=DiagnosisHint(
                failure_kind=str(d.get("failure_kind") or "unknown"),
                target_layer=str(d.get("target_layer") or "none"),
                expected_behavior=str(d.get("expected_behavior") or "")),
        ))
    return cases, problems, fixes


def mark_errors(path: str, problems: List[Dict], fixes: List[str], out_path: str = None) -> str:
    """Write problems back to the Excel: red-filled problem cells + a validation-result column and sheet."""
    from openpyxl import load_workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    out = out_path or path.replace(".xlsx", "") + "-需修改.xlsx"
    wb = load_workbook(path)
    ws = wb.active
    red = PatternFill(start_color=_BAD_FILL, end_color=_BAD_FILL, fill_type="solid")
    bad_font = Font(color=_BAD_FONT, bold=True)

    err_col = len(COLUMNS) + 1
    ws.cell(row=1, column=err_col, value="校验结果（改对后此列可删）").font = Font(bold=True, color=_BAD_FONT)

    by_row: Dict[int, List[str]] = {}
    for p in problems:
        by_row.setdefault(p["row"], []).append(p["msg"])
    for row_no, msgs in by_row.items():
        for ci in range(1, err_col + 1):
            ws.cell(row=row_no, column=ci).fill = red
        ws.cell(row=row_no, column=err_col, value="；".join(msgs)).font = bad_font
        ws.cell(row=row_no, column=err_col).fill = red

    if fixes:
        ws2 = wb.create_sheet("已自动补全")
        ws2.append(["项目"])
        for x in fixes:
            ws2.append([x])
    wb.save(out)
    return out
