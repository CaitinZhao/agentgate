"""W3 content-management endpoint tests: Excel round-trip (template -> fill -> import with
auto-fix and marked-error loop) and the guides endpoint."""
import openpyxl

from agentgate.webapp.routes.content_routes import _TMP_OK
from test_webapp import _seed_bank, _seed_users, platform


def test_excel_roundtrip(tmp_path, monkeypatch):
    with platform(tmp_path, monkeypatch) as client:
        users = _seed_users(client)
        _seed_bank(client, users["admin"])
        ah = users["admin"]
        # blank template downloads
        r = client.get("/api/v1/benchmarks/demo/cases/template?blank=true", headers=ah)
        assert r.status_code == 200 and r.content[:2] == b"PK"      # xlsx zip magic
        # build a filled sheet from the exported bank
        r = client.get("/api/v1/benchmarks/demo/cases/template", headers=ah)
        src = tmp_path / "export.xlsx"
        src.write_bytes(r.content)
        wb = openpyxl.load_workbook(src)
        ws = wb.active
        col = {c.value: i + 1 for i, c in enumerate(ws[1])}
        # fresh rows below the exported cases: one valid, one missing the query
        row = ws.max_row + 1
        ws.cell(row=row, column=col["case_id"], value="demo-xlsx1")
        ws.cell(row=row, column=col["query"], value="excel question")
        ws.cell(row=row, column=col["level"], value="L1")
        ws.cell(row=row, column=col["type(auto/numeric/boolean/extractive/free_text/refusal)"],
                value="numeric")
        ws.cell(row=row, column=col["gold_value"], value=42)
        ws.cell(row=row, column=col["gold_unit"], value="%")
        # a bad row: missing query
        ws.cell(row=row + 1, column=col["case_id"], value="demo-xlsx2")
        bad = tmp_path / "filled.xlsx"
        wb.save(bad)
        # first import: one problem row -> problems + marked download, nothing imported
        r = client.post("/api/v1/benchmarks/demo/cases/bulk", headers=ah,
                        files={"file": ("filled.xlsx", bad.read_bytes(),
                                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
        assert r.status_code == 200
        body = r.json()
        assert body["status"] == "problems" and len(body["problems"]) == 1
        assert body["marked_download"].startswith("/api/v1/tmp/")
        r = client.get(body["marked_download"], headers=ah)
        assert r.status_code == 200 and r.content[:2] == b"PK"
        # fix the sheet and re-import: the two new rows land
        wb2 = openpyxl.load_workbook(bad)
        ws2 = wb2.active
        ws2.cell(row=row + 1, column=col["query"], value="fixed question")
        fixed = tmp_path / "fixed.xlsx"
        wb2.save(fixed)
        r = client.post("/api/v1/benchmarks/demo/cases/bulk", headers=ah,
                        files={"file": ("fixed.xlsx", fixed.read_bytes(),
                                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
        assert r.status_code == 200 and r.json()["status"] == "ok"
        # the exported sheet carries the bank's 3 original cases + our 2 fresh rows
        assert r.json()["imported"] == 5
        # tmp-name whitelist rejects traversal (httpx normalizes ".." pre-routing -> 404)
        sc = client.get("/api/v1/tmp/..%2Fagentgate.db", headers=ah).status_code
        assert sc in (400, 404) and sc != 200


def test_guides_endpoint(tmp_path, monkeypatch):
    with platform(tmp_path, monkeypatch) as client:
        users = _seed_users(client)
        for guide in ("bank-edit", "bank-create", "run-eval"):
            r = client.get(f"/api/v1/guides/{guide}", headers=users["member"])
            assert r.status_code == 200 and len(r.text) > 200
            r = client.get(f"/api/v1/guides/{guide}?lang=en", headers=users["member"])
            assert r.status_code == 200
        assert client.get("/api/v1/guides/unknown", headers=users["member"]).status_code == 404
        # docs screenshots are served for the guide pages (extension-whitelisted, traversal-safe)
        ok = client.get("/api/v1/guides/images/06-run-create.png", headers=users["member"])
        assert ok.status_code == 200 and ok.content.startswith(bytes.fromhex("89504e47"))
        assert client.get("/api/v1/guides/images/nope.png", headers=users["member"]).status_code == 404
        assert client.get("/api/v1/guides/images/..%2F.env", headers=users["member"]).status_code in (400, 404)


def test_tmp_name_pattern():
    assert _TMP_OK.match("bulk-20260922120000-abc123-marked.xlsx")
    assert _TMP_OK.match("bulk-abc.jsonl")
    assert not _TMP_OK.match("../agentgate.db")
    assert not _TMP_OK.match("x.xlsx")
